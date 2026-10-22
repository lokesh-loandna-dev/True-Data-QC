import os
import sys
import re
import json
import argparse
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl
from sqlalchemy import text
from dotenv import load_dotenv

# Ensure backend root is on sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

load_dotenv(BASE_DIR / ".env")

from app.database.connection import init_db, engine
from app.routers.loans import (
    ensure_dict_cache,
    DATA_DICT,
    CONDITION_RULES,
    evaluate_rows,
    summarize,
    find_header_row,
    is_blank
)

# Default configuration
DEFAULT_SAMPLES_DIR = r"C:\Users\lsaravanan\Downloads\samples"
DEFAULT_USER_EMAIL = "lsaravanan@loandna.com"


def natural_sort_key(s):
    """Sort strings with embedded numbers naturally (e.g. Sample 1, Sample 2, Sample 10)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]


def get_user_by_email(conn, email: str):
    """Fetch user by email from the database."""
    query = text("SELECT id, username, email, role FROM users WHERE LOWER(email) = LOWER(:email)")
    return conn.execute(query, {"email": email.strip()}).mappings().first()


def list_all_users(conn):
    """Fetch all users to present as options when email is not found."""
    query = text("SELECT id, username, email, role FROM users ORDER BY id ASC")
    return conn.execute(query).mappings().all()


def generate_past_timestamps(count: int, base_time: datetime = None):
    """
    Generate distinct past dates and times specifically crafted to test:
    - Today preset
    - Yesterday preset
    - Last 7 days preset
    - Last 30 days preset
    - This Month preset
    - Custom date range filters
    All timestamps are strictly in the past.
    """
    if base_time is None:
        base_time = datetime.utcnow()

    # Pre-planned relative past offsets for up to 10 sample files
    preset_offsets = [
        {"delta": timedelta(hours=2, minutes=15), "label": "Today (~2 hrs ago)"},
        {"delta": timedelta(hours=6, minutes=45), "label": "Today (~7 hrs ago)"},
        {"delta": timedelta(days=1, hours=4, minutes=10), "label": "Yesterday"},
        {"delta": timedelta(days=2, hours=7, minutes=20), "label": "2 days ago (Last 7 Days)"},
        {"delta": timedelta(days=4, hours=3, minutes=45), "label": "4 days ago (Last 7 Days)"},
        {"delta": timedelta(days=6, hours=8, minutes=15), "label": "6 days ago (Last 7 Days)"},
        {"delta": timedelta(days=9, hours=5, minutes=30), "label": "9 days ago (Last 30 Days)"},
        {"delta": timedelta(days=14, hours=6, minutes=50), "label": "14 days ago (Last 30 Days)"},
        {"delta": timedelta(days=21, hours=4, minutes=25), "label": "21 days ago (Last 30 Days)"},
        {"delta": timedelta(days=28, hours=9, minutes=10), "label": "28 days ago (Last 30 Days)"},
    ]

    timestamps = []
    for idx in range(count):
        if idx < len(preset_offsets):
            dt = base_time - preset_offsets[idx]["delta"]
            label = preset_offsets[idx]["label"]
        else:
            # Fallback for more than 10 files
            days_ago = 30 + (idx - len(preset_offsets)) * 3
            dt = base_time - timedelta(days=days_ago, hours=5)
            label = f"{days_ago} days ago"
        
        # Round seconds to clean integers
        dt = dt.replace(microsecond=0)
        timestamps.append((dt, label))

    return timestamps


def process_and_ingest_excel(file_path: Path, target_timestamp: datetime, user_id: int, conn):
    """
    Validates and ingests a single Excel file using the exact business rules and schema
    as the True Data QC application.
    """
    filename = file_path.name

    # 1. Load Excel file
    workbook = openpyxl.load_workbook(str(file_path), data_only=True)
    worksheet = workbook.active

    # 2. Find header row
    header_row = find_header_row(worksheet)

    headers = {}
    excel_headers = []
    for cell in worksheet[header_row]:
        if cell.value is not None and str(cell.value).strip():
            h_name = str(cell.value).strip()
            headers.setdefault(h_name, cell.column)
            if h_name not in excel_headers:
                excel_headers.append(h_name)

    for required in ("Data Field", "Data Label", "True Data"):
        if required not in headers:
            raise ValueError(f"Missing required column '{required}' in {filename}. Found headers: {list(headers.keys())}")

    name_col = headers["Data Field"]
    label_col = headers["Data Label"]
    value_col = headers["True Data"]
    loan_id_col = headers.get("Loan ID")
    first_data_row = header_row + 1

    raw_field_names = [worksheet.cell(row=r, column=name_col).value
                       for r in range(first_data_row, worksheet.max_row + 1)]
    raw_labels = [worksheet.cell(row=r, column=label_col).value
                  for r in range(first_data_row, worksheet.max_row + 1)]
    raw_values = [worksheet.cell(row=r, column=value_col).value
                  for r in range(first_data_row, worksheet.max_row + 1)]
    raw_loan_ids = ([worksheet.cell(row=r, column=loan_id_col).value
                     for r in range(first_data_row, worksheet.max_row + 1)]
                    if loan_id_col else [None] * len(raw_labels))

    raw_all_rows = []
    for r in range(first_data_row, worksheet.max_row + 1):
        row_dict = {}
        for h_name in excel_headers:
            c_idx = headers[h_name]
            c_val = worksheet.cell(row=r, column=c_idx).value
            row_dict[h_name] = str(c_val) if c_val is not None else ""
        raw_all_rows.append(row_dict)

    # Filter out completely blank rows
    field_names = []
    labels = []
    values = []
    loan_ids = []
    filtered_all_rows = []
    for i in range(len(raw_labels)):
        if not (is_blank(raw_field_names[i]) and is_blank(raw_labels[i]) and is_blank(raw_values[i])):
            field_names.append(raw_field_names[i])
            labels.append(raw_labels[i])
            values.append(raw_values[i])
            loan_ids.append(raw_loan_ids[i])
            filtered_all_rows.append(raw_all_rows[i])

    if not labels:
        raise ValueError(f"{filename} does not contain any valid data rows.")

    # 3. Evaluate rows using Data Dictionary and Condition Rules
    ensure_dict_cache()
    from app.routers.loans import DATA_DICT
    results = evaluate_rows(
        field_names, labels, values, DATA_DICT, loan_ids=loan_ids,
        condition_rules=CONDITION_RULES
    )
    stats = summarize(filename, results, labels, DATA_DICT)

    # 4. Extract Loan ID
    extracted_loan_id = None
    for fn, val in zip(field_names, values):
        if fn and str(fn).strip().lower() in ("loan id", "loan_id", "loanid"):
            if val is not None and str(val).strip() != "":
                extracted_loan_id = str(val).strip()
                break

    if not extracted_loan_id:
        for lbl, val in zip(labels, values):
            if lbl and str(lbl).strip().lower() in ("loan id", "loan_id", "loanid"):
                if val is not None and str(val).strip() != "":
                    extracted_loan_id = str(val).strip()
                    break

    if not extracted_loan_id and any(lid is not None for lid in loan_ids):
        for lid in loan_ids:
            if lid is not None and str(lid).strip() != "":
                extracted_loan_id = str(lid).strip()
                break

    # 5. Insert Package into `loan_packages` with specific `uploaded_at` timestamp
    insert_package_sql = text("""
    INSERT INTO loan_packages (
        filename, loan_id, user_id, uploaded_at, total_fields, matched, unmatched, type_pass, type_fail, blank_values,
        total_enum, enum_pass, enum_format_variance, enum_fail, enum_no_list,
        numeric_format_flagged, total_macro, macro_valid, macro_invalid,
        final_valid, final_invalid, final_blank
    )
    OUTPUT INSERTED.id
    VALUES (
        :filename, :loan_id, :user_id, :uploaded_at, :total_fields, :matched, :unmatched, :type_pass, :type_fail, :blank_values,
        :total_enum, :enum_pass, :enum_format_variance, :enum_fail, :enum_no_list,
        :numeric_format_flagged, :total_macro, :macro_valid, :macro_invalid,
        :final_valid, :final_invalid, :final_blank
    )
    """)

    params = {
        "filename": filename,
        "loan_id": extracted_loan_id,
        "user_id": user_id,
        "uploaded_at": target_timestamp,
        "total_fields": int(stats.get("Total Fields", 0)),
        "matched": int(stats.get("Matched", 0)),
        "unmatched": int(stats.get("Unmatched", 0)),
        "type_pass": int(stats.get("Type Pass", 0)),
        "type_fail": int(stats.get("Type Fail", 0)),
        "blank_values": int(stats.get("Blank Values", 0)),
        "total_enum": int(stats.get("Total Enum", 0)),
        "enum_pass": int(stats.get("Enum Pass", 0)),
        "enum_format_variance": int(stats.get("Enum Format Variance", 0)),
        "enum_fail": int(stats.get("Enum Fail", 0)),
        "enum_no_list": int(stats.get("Enum No List", 0)),
        "numeric_format_flagged": int(stats.get("Numeric Format Flagged", 0)),
        "total_macro": int(stats.get("Total Macro", 0)),
        "macro_valid": int(stats.get("Macro Valid", 0)),
        "macro_invalid": int(stats.get("Macro Invalid", 0)),
        "final_valid": int(stats.get("Final Valid", 0)),
        "final_invalid": int(stats.get("Final Invalid", 0)),
        "final_blank": int(stats.get("Final Blank", 0))
    }

    package_id = conn.execute(insert_package_sql, params).scalar()

    # 6. Insert all records into `loan_package_records`
    insert_record_sql = text("""
    INSERT INTO loan_package_records (
        package_id, loan_id, data_field, data_label, true_data,
        dictionary_match_status, data_type_validation, data_type_remarks,
        numeric_format_flag, enum_validation, enum_remarks,
        macro_validation_result, result,
        final_result, final_remarks,
        condition_validation, condition_rule_type, condition_remarks,
        action, remark, raw_data
    )
    VALUES (
        :package_id, :loan_id, :data_field, :data_label, :true_data,
        :dictionary_match_status, :data_type_validation, :data_type_remarks,
        :numeric_format_flag, :enum_validation, :enum_remarks,
        :macro_validation_result, :result,
        :final_result, :final_remarks,
        :condition_validation, :condition_rule_type, :condition_remarks,
        :action, :remark, :raw_data
    )
    """)

    chunk_size = 1000
    records_to_insert = []
    for idx, row_res in enumerate(results):
        raw_row = filtered_all_rows[idx]
        records_to_insert.append({
            "package_id": package_id,
            "loan_id": str(loan_ids[idx]) if loan_ids[idx] is not None else (extracted_loan_id or None),
            "data_field": str(field_names[idx]) if field_names[idx] is not None else None,
            "data_label": str(labels[idx]) if labels[idx] is not None else None,
            "true_data": str(values[idx]) if values[idx] is not None else None,
            "dictionary_match_status": row_res[0],
            "data_type_validation": row_res[1],
            "data_type_remarks": row_res[2],
            "numeric_format_flag": row_res[3],
            "enum_validation": row_res[4],
            "enum_remarks": row_res[5],
            "macro_validation_result": row_res[6],
            "result": row_res[7],
            "final_result": row_res[11] if len(row_res) > 11 else row_res[7],
            "final_remarks": row_res[12] if len(row_res) > 12 else "",
            "condition_validation": row_res[8] if len(row_res) > 8 else None,
            "condition_rule_type": row_res[9] if len(row_res) > 9 else None,
            "condition_remarks": row_res[10] if len(row_res) > 10 else None,
            "action": None,
            "remark": None,
            "raw_data": json.dumps(raw_row) if raw_row else None
        })

    for i in range(0, len(records_to_insert), chunk_size):
        chunk = records_to_insert[i:i + chunk_size]
        conn.execute(insert_record_sql, chunk)

    return {
        "package_id": package_id,
        "filename": filename,
        "loan_id": extracted_loan_id,
        "uploaded_at": target_timestamp,
        "total_fields": stats.get("Total Fields", 0),
        "final_valid": stats.get("Final Valid", 0),
        "final_invalid": stats.get("Final Invalid", 0)
    }


def main():
    parser = argparse.ArgumentParser(
        description="Ingest sample loan Excel files with varied past dates and times for testing dashboard date filters."
    )
    parser.add_argument(
        "--samples-dir",
        default=DEFAULT_SAMPLES_DIR,
        help=f"Directory containing sample Excel files (default: {DEFAULT_SAMPLES_DIR})"
    )
    parser.add_argument(
        "--email",
        default=DEFAULT_USER_EMAIL,
        help=f"Email of the user account under which to ingest the packages (default: {DEFAULT_USER_EMAIL})"
    )
    parser.add_argument(
        "--clear-existing",
        action="store_true",
        help="Clear existing loan packages before ingesting new sample data"
    )

    args = parser.parse_args()

    print("=" * 80)
    print(" TRUE DATA QC - SAMPLE INGESTION WITH PAST DATES & USER MAPPING")
    print("=" * 80)

    # 1. Initialize DB
    init_db()
    if not engine:
        print("[!] Error: Could not connect to the database. Check .env DATABASE_URL.")
        sys.exit(1)

    # 2. Check and locate user by email
    with engine.connect() as conn:
        target_user = get_user_by_email(conn, args.email)
        if not target_user:
            print(f"[x] User with email '{args.email}' was NOT found in the database!")
            print("\nAvailable registered users in database:")
            users = list_all_users(conn)
            for u in users:
                print(f"  - ID: {u['id']}, Username: {u['username']}, Email: {u['email']}, Role: {u['role']}")
            print(f"\nPlease re-run the script with --email <one_of_the_above_emails>")
            sys.exit(1)

        print(f"[+] User Verified for Ingestion:")
        print(f"    - ID:       {target_user['id']}")
        print(f"    - Username: {target_user['username']}")
        print(f"    - Email:    {target_user['email']}")
        print(f"    - Role:     {target_user['role']}")

        # Optional: clear existing packages
        if args.clear_existing:
            print("\n[-] Clearing existing loan packages as requested (--clear-existing)...")
            conn.execute(text("DELETE FROM loan_package_records"))
            conn.execute(text("DELETE FROM loan_packages"))
            conn.commit()
            print("[+] Previous packages cleared.")

    # 3. Locate sample files
    samples_dir = Path(args.samples_dir)
    if not samples_dir.exists():
        print(f"[!] Error: Samples directory does not exist: {samples_dir}")
        sys.exit(1)

    excel_files = sorted(
        [f for f in samples_dir.glob("*.xlsx") if not f.name.startswith("~$")],
        key=natural_sort_key
    )

    if not excel_files:
        print(f"[!] No Excel (.xlsx) files found in {samples_dir}")
        sys.exit(1)

    print(f"\n[+] Found {len(excel_files)} Excel file(s) in: {samples_dir}")

    # 4. Generate past timestamps
    timestamps = generate_past_timestamps(len(excel_files))

    print("\n[+] Planned Ingestion Timeline (Past Dates Only):")
    print("-" * 80)
    for i, (f, (ts, label)) in enumerate(zip(excel_files, timestamps), start=1):
        print(f" {i:2d}. {f.name:<20} -> {ts.strftime('%Y-%m-%d %H:%M:%S')} ({label})")
    print("-" * 80)

    # 5. Ingest each file
    print(f"\n[+] Ingesting {len(excel_files)} loan package(s) into database...")
    ingested_summary = []

    with engine.connect() as conn:
        for idx, (f_path, (ts, label)) in enumerate(zip(excel_files, timestamps), start=1):
            print(f"\n[{idx}/{len(excel_files)}] Processing: {f_path.name}...")
            try:
                res = process_and_ingest_excel(f_path, ts, target_user["id"], conn)
                conn.commit()
                ingested_summary.append({**res, "label": label})
                print(f"    [OK] Package ID: {res['package_id']} | Loan ID: {res['loan_id']} | Date: {ts} | Fields: {res['total_fields']}")
            except Exception as e:
                conn.rollback()
                print(f"    [ERROR] Failed to ingest {f_path.name}: {e}")

    # 6. Final report
    print("\n" + "=" * 80)
    print(f" INGESTION COMPLETE: {len(ingested_summary)} / {len(excel_files)} PACKAGES INGESTED")
    print(f" Ingested User: {target_user['username']} ({target_user['email']})")
    print("=" * 80)
    print(f"{'Pkg ID':<8} {'Filename':<18} {'Loan ID':<16} {'Ingested Date/Time':<22} {'Test Preset'}")
    print("-" * 80)
    for row in ingested_summary:
        dt_str = row['uploaded_at'].strftime('%Y-%m-%d %H:%M:%S')
        print(f"{row['package_id']:<8} {row['filename']:<18} {str(row['loan_id']):<16} {dt_str:<22} {row['label']}")
    print("-" * 80)
    print("\nYou can now open the Dashboard in your browser and test:")
    print("  - 'Today' filter (captures packages ingested today)")
    print("  - 'Yesterday' filter (captures yesterday's packages)")
    print("  - 'Last 7 Days' filter (captures packages from the past week)")
    print("  - 'Last 30 Days' filter (captures all past packages within 30 days)")
    print("  - Custom 'From Date' & 'To Date' pickers")


if __name__ == "__main__":
    main()
