import io
import json
import random
import re
import time
from datetime import datetime
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from fastapi.responses import StreamingResponse
from app.auth import get_current_user_id
from pydantic import BaseModel
from typing import List, Dict, Any, Optional
from sqlalchemy import text
import pandas as pd
import openpyxl

from pathlib import Path
from app.database.connection import engine
from lm_helper.validate_loan_files_combined_v3 import evaluate_rows, summarize, find_header_row, is_blank
from lm_helper import condition_validation as cv

router = APIRouter(
    prefix="/api/loans",
    tags=["Loans"]
)

def _is_deadlock(exc: Exception) -> bool:
    s = str(exc)
    return "1205" in s or "deadlock" in s.lower()

def _retry_on_deadlock(fn, max_retries: int = 3):
    for attempt in range(max_retries):
        try:
            return fn()
        except Exception as exc:
            if attempt < max_retries - 1 and _is_deadlock(exc):
                time.sleep(0.05 * (2 ** attempt) + random.random() * 0.05)
                continue
            raise

# Load condition rules once at startup if spec exists
CONDITION_SPEC_PATH = Path(__file__).resolve().parent.parent.parent / "lm_helper" / "Condition_Validation_v2.xlsx"
CONDITION_RULES = None
if CONDITION_SPEC_PATH.exists():
    try:
        CONDITION_RULES = cv.load_rules(CONDITION_SPEC_PATH)
    except Exception as e:
        print(f"Warning: Failed to load Condition_Validation spec: {e}")

# Load data dictionary once at startup — same pattern as CONDITION_RULES above.
# Eliminates SELECT * FROM data_dictionary (~6,962 rows) on every /validate request.
# Call reload_dict_cache() after a dictionary upload to refresh without restarting.
_LOCAL_DICT_VERSION = None
DATA_DICT = None

def _format_utc_iso(dt_val):
    """Serialize a database datetime (stored in UTC) as an ISO-8601 string with 'Z' suffix."""
    if not dt_val:
        return None
    if isinstance(dt_val, str):
        if not dt_val.endswith("Z") and "+" not in dt_val:
            return dt_val + "Z"
        return dt_val
    iso = dt_val.isoformat()
    if not iso.endswith("Z") and "+" not in iso:
        return iso + "Z"
    return iso

def _build_dict_df(rows):
    data_list = []
    for row in rows:
        r = dict(row)
        fn = (
            r.get("field_name") or r.get("Field Name") or r.get("fieldname")
            or r.get("data_field") or r.get("Data Field") or r.get("field") or ""
        )
        fl = (
            r.get("field_label") or r.get("Field Label") or r.get("fieldlabel")
            or r.get("data_label") or r.get("Data Label") or r.get("label") or ""
        )
        dt = (
            r.get("data_type") or r.get("Data Type") or r.get("datatype")
            or r.get("type") or ""
        )
        fe = (
            r.get("field_enumerations") or r.get("Field Enumerations") or r.get("fieldenumerations")
            or r.get("field_enumeration") or r.get("enumerations") or r.get("enum") or ""
        )
        data_list.append({
            "Field Name": str(fn) if fn is not None else "",
            "Field Label": str(fl) if fl is not None else "",
            "Data Type": str(dt) if dt is not None else "",
            "Field Enumerations": str(fe) if fe is not None else ""
        })
    df = pd.DataFrame(data_list)
    if df.empty:
        return pd.DataFrame(columns=["Field Name", "Field Label", "Data Type", "Field Enumerations"])
    # Filter out empty or whitespace-only labels
    df = df[df["Field Label"].astype(str).str.strip() != ""].copy()
    df["_key"] = df["Field Label"].astype(str).str.strip()
    df = df.drop_duplicates("_key")
    return df.set_index("_key")

def reload_dict_cache(version=None):
    global DATA_DICT, _LOCAL_DICT_VERSION
    if not engine:
        return
    try:
        from app.database.operations import get_data_dict_version
        v = version or get_data_dict_version()
        with engine.connect() as conn:
            table_exists = conn.execute(text(
                "SELECT COUNT(*) FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME = 'data_dictionary'"
            )).scalar()
            if table_exists:
                rows = conn.execute(text("SELECT * FROM dbo.data_dictionary")).mappings().all()
                DATA_DICT = _build_dict_df(rows)
            else:
                DATA_DICT = _build_dict_df([])
        _LOCAL_DICT_VERSION = v
    except Exception as e:
        print(f"Warning: Failed to load data dictionary: {e}")

def ensure_dict_cache():
    global DATA_DICT, _LOCAL_DICT_VERSION
    try:
        from app.database.operations import get_data_dict_version
        db_version = get_data_dict_version()
        if _LOCAL_DICT_VERSION != db_version or DATA_DICT is None:
            reload_dict_cache(db_version)
    except Exception as e:
        print(f"Warning: Failed to check dict cache version: {e}")

try:
    ensure_dict_cache()
except Exception:
    pass


class LoanRecordInput(BaseModel):
    loan_id: Optional[str] = None
    data_field: Optional[str] = None
    data_label: Optional[str] = None
    true_data: Optional[str] = None
    dictionary_match_status: Optional[str] = None
    data_type_validation: Optional[str] = None
    data_type_remarks: Optional[str] = None
    numeric_format_flag: Optional[str] = None
    enum_validation: Optional[str] = None
    enum_remarks: Optional[str] = None
    macro_validation_result: Optional[str] = None
    result: Optional[str] = None
    condition_validation: Optional[str] = None
    condition_rule_type: Optional[str] = None
    condition_remarks: Optional[str] = None
    condition_text: Optional[str] = None
    condition_outcome: Optional[str] = None
    final_result: Optional[str] = None
    final_remarks: Optional[str] = None
    action: Optional[str] = None
    remark: Optional[str] = None
    raw_data: Optional[Dict[str, Any]] = None

class LoanSaveRequest(BaseModel):
    filename: str
    summary: Dict[str, Any]
    records: List[LoanRecordInput]
    loan_id: Optional[str] = None

@router.post("/validate")
def validate_loan_file(
    file: UploadFile = File(...),
    user_id: int = Depends(get_current_user_id)
):
    if not file.filename.endswith(('.xlsx', '.xls', '.xlsm')):
        raise HTTPException(status_code=400, detail="Only Excel files (.xlsx, .xls, .xlsm) are supported")

    try:
        contents = file.file.read()

        ensure_dict_cache()
        dictionary = DATA_DICT
        if dictionary is None or len(dictionary) == 0:
            raise HTTPException(
                status_code=400,
                detail="Data dictionary is empty in the database. Please upload a valid Data Dictionary Excel file in the Data Dictionary tab first before validating loan files."
            )

        # Load uploaded Excel
        workbook = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        worksheet = workbook.active

        try:
            header_row = find_header_row(worksheet)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))

        headers = {}
        excel_headers = []
        for cell in worksheet[header_row]:
            if cell.value is not None and str(cell.value).strip():
                h_name = str(cell.value).strip()
                headers.setdefault(h_name, cell.column)
                if h_name not in excel_headers:
                    excel_headers.append(h_name)

        for required in ("Data Field", "Data Label", "True Data"):
            if required not in headers:
                raise HTTPException(
                    status_code=400,
                    detail=f"Missing required column '{required}' in sheet '{worksheet.title}'. Found headers: {list(headers.keys())}"
                )

        name_col = headers["Data Field"]
        label_col = headers["Data Label"]
        value_col = headers["True Data"]
        loan_id_col = headers.get("Loan ID")
        first_data_row = header_row + 1

        raw_field_names = [worksheet.cell(row=r, column=name_col).value
                           for r in range(first_data_row, worksheet.max_row + 1)]
        raw_labels = [worksheet.cell(row=r, column=label_col).value
                      for r in range(first_data_row, worksheet.max_row + 1)]
        raw_values = [worksheet.cell(row=r, column=value_col).value
                      for r in range(first_data_row, worksheet.max_row + 1)]
        raw_loan_ids = ([worksheet.cell(row=r, column=loan_id_col).value
                         for r in range(first_data_row, worksheet.max_row + 1)]
                        if loan_id_col else [None] * len(raw_labels))

        # Capture all row values for all headers present in Excel
        raw_all_rows = []
        for r in range(first_data_row, worksheet.max_row + 1):
            row_dict = {}
            for h_name in excel_headers:
                c_idx = headers[h_name]
                c_val = worksheet.cell(row=r, column=c_idx).value
                row_dict[h_name] = str(c_val) if c_val is not None else ""
            raw_all_rows.append(row_dict)

        # Filter out completely blank rows
        field_names = []
        labels = []
        values = []
        loan_ids = []
        filtered_all_rows = []
        for i in range(len(raw_labels)):
            if not (is_blank(raw_field_names[i]) and is_blank(raw_labels[i]) and is_blank(raw_values[i])):
                field_names.append(raw_field_names[i])
                labels.append(raw_labels[i])
                values.append(raw_values[i])
                loan_ids.append(raw_loan_ids[i])
                filtered_all_rows.append(raw_all_rows[i])

        if not labels:
            raise HTTPException(status_code=400, detail="The Excel sheet does not contain any valid data rows.")

        # Run validation
        results = evaluate_rows(
            field_names, labels, values, dictionary, loan_ids=loan_ids,
            condition_rules=CONDITION_RULES
        )
        stats = summarize(file.filename, results, labels, dictionary)
        if "Review Required" not in stats and "Overall Review Required" in stats:
            stats["Review Required"] = stats["Overall Review Required"]

        # Build response records list with all excel fields included
        response_records = []
        for idx, row_res in enumerate(results):
            raw_row = filtered_all_rows[idx]
            has_condition_text_outcome = len(row_res) >= 15
            cond_val = row_res[8] if len(row_res) > 8 else None
            cond_type = row_res[9] if len(row_res) > 9 else None
            cond_rem = row_res[10] if len(row_res) > 10 else None
            cond_text = row_res[11] if has_condition_text_outcome else None
            cond_outcome = row_res[12] if has_condition_text_outcome else None
            fin_res = row_res[13] if has_condition_text_outcome else (row_res[11] if len(row_res) > 11 else row_res[7])
            fin_rem = row_res[14] if has_condition_text_outcome else (row_res[12] if len(row_res) > 12 else "")

            # Ensure Condition Text and Condition Outcome are in raw_row as well
            if cond_text is not None:
                raw_row["Condition Text"] = cond_text
            if cond_outcome is not None:
                raw_row["Condition Outcome"] = cond_outcome

            rec = {
                "loan_id": str(loan_ids[idx]) if loan_ids[idx] is not None else None,
                "data_field": str(field_names[idx]) if field_names[idx] is not None else None,
                "data_label": str(labels[idx]) if labels[idx] is not None else None,
                "true_data": str(values[idx]) if values[idx] is not None else None,
                "dictionary_match_status": row_res[0],
                "data_type_validation": row_res[1],
                "data_type_remarks": row_res[2],
                "numeric_format_flag": row_res[3],
                "enum_validation": row_res[4],
                "enum_remarks": row_res[5],
                "macro_validation_result": row_res[6],
                "result": row_res[7],
                "condition_validation": cond_val,
                "condition_rule_type": cond_type,
                "condition_remarks": cond_rem,
                "condition_text": cond_text,
                "condition_outcome": cond_outcome,
                "final_result": fin_res,
                "final_remarks": fin_rem,
                "action": "",
                "remark": "",
                "raw_data": raw_row
            }
            # Merge raw_row keys into record if key not already in rec
            for k, v in raw_row.items():
                if k not in rec:
                    rec[k] = v
            response_records.append(rec)

        # Extract loan_id from row where Data Field value is 'Loan ID' (True Data value)
        extracted_loan_id = None
        for fn, val in zip(field_names, values):
            if fn and str(fn).strip().lower() in ("loan id", "loan_id", "loanid"):
                if val is not None and str(val).strip() != "":
                    extracted_loan_id = str(val).strip()
                    break

        if not extracted_loan_id:
            for lbl, val in zip(labels, values):
                if lbl and str(lbl).strip().lower() in ("loan id", "loan_id", "loanid"):
                    if val is not None and str(val).strip() != "":
                        extracted_loan_id = str(val).strip()
                        break

        if not extracted_loan_id and any(lid is not None for lid in loan_ids):
            for lid in loan_ids:
                if lid is not None and str(lid).strip() != "":
                    extracted_loan_id = str(lid).strip()
                    break

        return {
            "filename": file.filename,
            "loan_id": extracted_loan_id,
            "headers": excel_headers,
            "summary": stats,
            "records": response_records
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/save")
def save_loan_package(body: LoanSaveRequest, user_id: int = Depends(get_current_user_id)):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    try:
        # Determine loan_id for this package
        loan_id_val = body.loan_id
        if not loan_id_val:
            for r in body.records:
                df_clean = str(r.data_field or '').strip().lower()
                if df_clean in ("loan id", "loan_id", "loanid"):
                    if r.true_data is not None and str(r.true_data).strip() != "":
                        loan_id_val = str(r.true_data).strip()
                        break
        if not loan_id_val:
            for r in body.records:
                dl_clean = str(r.data_label or '').strip().lower()
                if dl_clean in ("loan id", "loan_id", "loanid"):
                    if r.true_data is not None and str(r.true_data).strip() != "":
                        loan_id_val = str(r.true_data).strip()
                        break
        if not loan_id_val:
            for r in body.records:
                if r.loan_id is not None and str(r.loan_id).strip() != "":
                    loan_id_val = str(r.loan_id).strip()
                    break

        insert_package_sql = text("""
        INSERT INTO loan_packages (
            filename, loan_id, user_id, uploaded_at, total_fields, matched, unmatched, type_pass, type_fail, blank_values,
            total_enum, enum_pass, enum_format_variance, enum_fail, enum_no_list,
            numeric_format_flagged, total_macro, macro_valid, macro_invalid,
            final_valid, final_invalid, final_blank, review_required
        )
        OUTPUT INSERTED.id
        VALUES (
            :filename, :loan_id, :user_id, SYSUTCDATETIME(), :total_fields, :matched, :unmatched, :type_pass, :type_fail, :blank_values,
            :total_enum, :enum_pass, :enum_format_variance, :enum_fail, :enum_no_list,
            :numeric_format_flagged, :total_macro, :macro_valid, :macro_invalid,
            :final_valid, :final_invalid, :final_blank, :review_required
        )
        """)

        summary_dict = body.summary
        params = {
            "filename": body.filename,
            "loan_id": loan_id_val,
            "user_id": user_id,
            "total_fields": int(summary_dict.get("Total Fields", 0)),
            "matched": int(summary_dict.get("Matched", 0)),
            "unmatched": int(summary_dict.get("Unmatched", 0)),
            "type_pass": int(summary_dict.get("Type Pass", 0)),
            "type_fail": int(summary_dict.get("Type Fail", 0)),
            "blank_values": int(summary_dict.get("Blank Values", 0)),
            "total_enum": int(summary_dict.get("Total Enum", 0)),
            "enum_pass": int(summary_dict.get("Enum Pass", 0)),
            "enum_format_variance": int(summary_dict.get("Enum Format Variance", 0)),
            "enum_fail": int(summary_dict.get("Enum Fail", 0)),
            "enum_no_list": int(summary_dict.get("Enum No List", 0)),
            "numeric_format_flagged": int(summary_dict.get("Numeric Format Flagged", 0)),
            "total_macro": int(summary_dict.get("Total Macro", 0)),
            "macro_valid": int(summary_dict.get("Macro Valid", 0)),
            "macro_invalid": int(summary_dict.get("Macro Invalid", 0)),
            "final_valid": int(summary_dict.get("Final Valid", 0)),
            "final_invalid": int(summary_dict.get("Final Invalid", 0)),
            "final_blank": int(summary_dict.get("Final Blank", 0)),
            "review_required": int(summary_dict.get("Review Required", 0) or summary_dict.get("Overall Review Required", 0))
        }

        insert_record_sql = text("""
        INSERT INTO loan_package_records (
            package_id, loan_id, data_field, data_label, true_data,
            dictionary_match_status, data_type_validation, data_type_remarks,
            numeric_format_flag, enum_validation, enum_remarks,
            macro_validation_result, result,
            final_result, final_remarks,
            condition_validation, condition_rule_type, condition_remarks,
            action, remark, raw_data
        )
        VALUES (
            :package_id, :loan_id, :data_field, :data_label, :true_data,
            :dictionary_match_status, :data_type_validation, :data_type_remarks,
            :numeric_format_flag, :enum_validation, :enum_remarks,
            :macro_validation_result, :result,
            :final_result, :final_remarks,
            :condition_validation, :condition_rule_type, :condition_remarks,
            :action, :remark, :raw_data
        )
        """)

        def _do_save():
            with engine.begin() as conn:
                pkg_id = conn.execute(insert_package_sql, params).scalar()

                # Bulk insert records in chunks
                chunk_size = 1000
                for i in range(0, len(body.records), chunk_size):
                    chunk = body.records[i:i + chunk_size]
                    record_params = []
                    for r in chunk:
                        raw_dict = dict(r.raw_data) if isinstance(r.raw_data, dict) else {}
                        if r.condition_text is not None and "Condition Text" not in raw_dict:
                            raw_dict["Condition Text"] = r.condition_text
                        if r.condition_outcome is not None and "Condition Outcome" not in raw_dict:
                            raw_dict["Condition Outcome"] = r.condition_outcome
                        raw_dict.pop("condition_text", None)
                        raw_dict.pop("condition_outcome", None)

                        record_params.append({
                            "package_id": pkg_id,
                            "loan_id": str(r.loan_id) if r.loan_id is not None else None,
                            "data_field": str(r.data_field) if r.data_field is not None else None,
                            "data_label": str(r.data_label) if r.data_label is not None else None,
                            "true_data": str(r.true_data) if r.true_data is not None else None,
                            "dictionary_match_status": str(r.dictionary_match_status) if r.dictionary_match_status is not None else None,
                            "data_type_validation": str(r.data_type_validation) if r.data_type_validation is not None else None,
                            "data_type_remarks": str(r.data_type_remarks) if r.data_type_remarks is not None else None,
                            "numeric_format_flag": str(r.numeric_format_flag) if r.numeric_format_flag is not None else None,
                            "enum_validation": str(r.enum_validation) if r.enum_validation is not None else None,
                            "enum_remarks": str(r.enum_remarks) if r.enum_remarks is not None else None,
                            "macro_validation_result": str(r.macro_validation_result) if r.macro_validation_result is not None else None,
                            "result": str(r.result) if r.result is not None else None,
                            "final_result": str(r.final_result) if r.final_result is not None else (str(r.result) if r.result is not None else None),
                            "final_remarks": str(r.final_remarks) if r.final_remarks is not None else None,
                            "condition_validation": str(r.condition_validation) if r.condition_validation is not None else None,
                            "condition_rule_type": str(r.condition_rule_type) if r.condition_rule_type is not None else None,
                            "condition_remarks": str(r.condition_remarks) if r.condition_remarks is not None else None,
                            "action": str(r.action) if r.action is not None else None,
                            "remark": str(r.remark) if r.remark is not None else None,
                            "raw_data": json.dumps(raw_dict) if raw_dict else None
                        })
                    conn.execute(insert_record_sql, record_params)
                return pkg_id

        package_id = _retry_on_deadlock(_do_save)

        # Build response from request body — avoids a full re-SELECT of all inserted rows
        saved_records = []
        for r in body.records:
            rd = {
                "package_id": package_id,
                "loan_id": str(r.loan_id) if r.loan_id is not None else None,
                "data_field": str(r.data_field) if r.data_field is not None else None,
                "data_label": str(r.data_label) if r.data_label is not None else None,
                "true_data": str(r.true_data) if r.true_data is not None else None,
                "dictionary_match_status": str(r.dictionary_match_status) if r.dictionary_match_status is not None else None,
                "data_type_validation": str(r.data_type_validation) if r.data_type_validation is not None else None,
                "data_type_remarks": str(r.data_type_remarks) if r.data_type_remarks is not None else None,
                "numeric_format_flag": str(r.numeric_format_flag) if r.numeric_format_flag is not None else None,
                "enum_validation": str(r.enum_validation) if r.enum_validation is not None else None,
                "enum_remarks": str(r.enum_remarks) if r.enum_remarks is not None else None,
                "macro_validation_result": str(r.macro_validation_result) if r.macro_validation_result is not None else None,
                "result": str(r.result) if r.result is not None else None,
                "final_result": str(r.final_result) if r.final_result is not None else (str(r.result) if r.result is not None else None),
                "final_remarks": str(r.final_remarks) if r.final_remarks is not None else None,
                "condition_validation": str(r.condition_validation) if r.condition_validation is not None else None,
                "condition_rule_type": str(r.condition_rule_type) if r.condition_rule_type is not None else None,
                "condition_remarks": str(r.condition_remarks) if r.condition_remarks is not None else None,
                "action": str(r.action) if r.action is not None else None,
                "remark": str(r.remark) if r.remark is not None else None,
                "raw_data": r.raw_data,
            }
            if isinstance(r.raw_data, dict):
                for k, v in r.raw_data.items():
                    if k not in rd:
                        rd[k] = v
            saved_records.append(rd)

        return {"status": "success", "package_id": package_id, "records": saved_records}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/")
def get_saved_packages(user_id: int = Depends(get_current_user_id)):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    try:
        with engine.connect() as conn:
            user_row = conn.execute(
                text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
            ).fetchone()
            if not user_row:
                raise HTTPException(status_code=404, detail="User not found")
            role = user_row[0]

            query = """
            SELECT 
                lp.*, 
                u.username as uploaded_by, 
                u.email as uploaded_email,
                COALESCE(fr.fr_valid, lp.final_valid) as final_valid,
                COALESCE(fr.fr_invalid, lp.final_invalid) as final_invalid,
                COALESCE(fr.fr_review_required, 0) as review_required,
                COALESCE(act.action_corrected, 0) as action_corrected,
                COALESCE(act.action_incorrect_status, 0) as action_incorrect_status,
                COALESCE(act.action_not_applicable, 0) as action_not_applicable,
                COALESCE(act.action_need_review, 0) as action_need_review,
                COALESCE(act.action_review_completed, 0) as action_review_completed,
                COALESCE(act.action_out_of_scope, 0) as action_out_of_scope,
                COALESCE(act.action_others, 0) as action_others,
                COALESCE(act.yet_to_take_action, 0) as yet_to_take_action
            FROM loan_packages lp
            LEFT JOIN users u ON lp.user_id = u.id
            LEFT JOIN (
                SELECT 
                    package_id,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(final_result))) = 'valid' THEN 1 ELSE 0 END) as fr_valid,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(final_result))) = 'invalid' THEN 1 ELSE 0 END) as fr_invalid,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(final_result))) = 'review required' THEN 1 ELSE 0 END) as fr_review_required
                FROM loan_package_records
                GROUP BY package_id
            ) fr ON lp.id = fr.package_id
            LEFT JOIN (
                SELECT 
                    package_id,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'corrected' THEN 1 ELSE 0 END) as action_corrected,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'incorrect status' THEN 1 ELSE 0 END) as action_incorrect_status,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'not applicable' THEN 1 ELSE 0 END) as action_not_applicable,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'need review' THEN 1 ELSE 0 END) as action_need_review,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'review completed' THEN 1 ELSE 0 END) as action_review_completed,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'out of scope' THEN 1 ELSE 0 END) as action_out_of_scope,
                    SUM(CASE WHEN LOWER(LTRIM(RTRIM(action))) = 'others' THEN 1 ELSE 0 END) as action_others,
                    SUM(CASE 
                        WHEN LOWER(LTRIM(RTRIM(final_result))) IN ('invalid', 'review required')
                         AND (action IS NULL OR LTRIM(RTRIM(action)) = '' OR LOWER(LTRIM(RTRIM(action))) IN ('select', '-- select --', 'none'))
                        THEN 1 ELSE 0 
                    END) as yet_to_take_action
                FROM loan_package_records
                GROUP BY package_id
            ) act ON lp.id = act.package_id
            """
            params = {}
            if role != "admin":
                query += " WHERE lp.user_id = :user_id"
                params["user_id"] = user_id
            query += " ORDER BY lp.uploaded_at DESC"

            result = conn.execute(text(query), params).mappings().all()
            packages = []
            for row in result:
                p = dict(row)
                if p.get("uploaded_at"):
                    p["uploaded_at"] = _format_utc_iso(p["uploaded_at"])
                packages.append(p)

        return {"status": "success", "packages": packages}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

class BulkExportRequest(BaseModel):
    package_ids: List[int]

@router.post("/bulk-export")
def bulk_export_packages(
    req: BulkExportRequest,
    user_id: int = Depends(get_current_user_id)
):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    if not req.package_ids:
        raise HTTPException(status_code=400, detail="Please select at least one package to export.")

    clean_ids = [int(pid) for pid in req.package_ids if pid is not None]
    if not clean_ids:
        raise HTTPException(status_code=400, detail="Invalid package IDs provided.")

    id_clause = ",".join(str(pid) for pid in clean_ids)

    try:
        with engine.connect() as conn:
            user_row = conn.execute(
                text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
            ).fetchone()
            if not user_row:
                raise HTTPException(status_code=404, detail="User not found")
            role = user_row[0]

            auth_filter = f"WHERE lp.id IN ({id_clause})"
            if role != "admin":
                auth_filter += f" AND lp.user_id = {int(user_id)}"

            packages_rows = conn.execute(text(f"""
                SELECT lp.id, lp.filename, lp.loan_id, lp.user_id,
                       u.username AS uploaded_by, u.email AS uploaded_email
                FROM loan_packages lp
                LEFT JOIN users u ON lp.user_id = u.id
                {auth_filter}
                ORDER BY lp.id DESC
            """)).mappings().all()

            if not packages_rows:
                raise HTTPException(status_code=404, detail="No authorized packages found for export.")

            packages_dict = {p["id"]: dict(p) for p in packages_rows}
            valid_ids_clause = ",".join(str(p["id"]) for p in packages_rows)

            records_rows = conn.execute(text(f"""
                SELECT * FROM loan_package_records
                WHERE package_id IN ({valid_ids_clause})
                ORDER BY package_id DESC, id ASC
            """)).mappings().all()

        # Extract dynamic raw_data columns and parse records
        dynamic_cols_set = set()
        dynamic_cols_order = []
        parsed_records = []

        standard_fixed_cols = {
            "File Name", "Data Field", "Data Label", "True Data",
            "Loan ID", "loan_id", "Loan Id", "Username", "Email",
            "Dict Match", "Type Check", "Type Remarks", "Numeric Format",
            "Enum Check", "Enum Remarks", "Macro Check", "Combined Result",
            "Condition Check", "Condition Remarks", "Condition Text", "Condition Outcome",
            "Final Result", "Final Remarks",
            "Action", "User Remark"
        }

        for r in records_rows:
            rec = dict(r)
            pkg = packages_dict.get(rec.get("package_id")) or {}
            pkg_filename = pkg.get("filename") or "Unknown"

            raw_dict = {}
            raw_str = rec.get("raw_data")
            if raw_str:
                try:
                    raw_dict = json.loads(raw_str) if isinstance(raw_str, str) else raw_dict
                except Exception:
                    raw_dict = {}

            if isinstance(raw_dict, dict):
                for k in raw_dict.keys():
                    if k not in standard_fixed_cols and k not in dynamic_cols_set:
                        dynamic_cols_set.add(k)
                        dynamic_cols_order.append(k)

            parsed_records.append((rec, pkg, pkg_filename, raw_dict))

        # Assemble final ordered headers:
        # 1. 'File Name'
        # 2. Dynamic raw_data columns
        # 3. Standard data columns ('Data Field', 'Data Label', 'True Data')
        # 4. Package Info columns ('Loan Id', 'Username', 'Email')
        # 5. Audit & Validation outcome columns
        final_headers = ["File Name"]
        for c in dynamic_cols_order:
            if c not in final_headers:
                final_headers.append(c)

        for std_col in ["Data Field", "Data Label", "True Data", "Loan Id", "Username", "Email"]:
            if std_col not in final_headers:
                final_headers.append(std_col)

        audit_cols = [
            "Dict Match", "Type Check", "Type Remarks", "Numeric Format",
            "Enum Check", "Enum Remarks", "Macro Check", "Combined Result",
            "Condition Check", "Condition Remarks", "Condition Text", "Condition Outcome",
            "Final Result", "Final Remarks",
            "Action", "User Remark"
        ]
        for ac in audit_cols:
            if ac not in final_headers:
                final_headers.append(ac)

        # Build workbook with openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Result Sheet"

        # Write header row
        ws.append(final_headers)

        # Style header row
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        header_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="94A3B8")
        )

        for col_idx in range(1, len(final_headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        # Populate rows
        for rec, pkg, pkg_filename, raw_dict in parsed_records:
            effective_loan_id = pkg.get("loan_id") or ""
            if not effective_loan_id and pkg_filename:
                m = re.search(r"(?:loan\s*id|loanid)[^\d]*(\d+)", pkg_filename, re.I) or re.search(r"\b(\d{7,12})\b", pkg_filename)
                if m:
                    effective_loan_id = m.group(1)

            row_obj = {"File Name": pkg_filename}
            if isinstance(raw_dict, dict):
                for k, v in raw_dict.items():
                    row_obj[k] = v if v is not None else ""

            if "Data Field" not in row_obj or not row_obj["Data Field"]:
                row_obj["Data Field"] = rec.get("data_field") or ""
            if "Data Label" not in row_obj or not row_obj["Data Label"]:
                row_obj["Data Label"] = rec.get("data_label") or ""
            if "True Data" not in row_obj or not row_obj["True Data"]:
                row_obj["True Data"] = rec.get("true_data") or ""

            # Remove duplicate loan id keys
            row_obj.pop("Loan ID", None)
            row_obj.pop("loan_id", None)

            # Resolve loan id, username, email
            row_loan_id = (
                rec.get("pkg_loan_id")
                or (str(rec.get("loan_id")).strip() if rec.get("loan_id") is not None and str(rec.get("loan_id")).strip() != "" else None)
                or (raw_dict.get("Loan ID") or raw_dict.get("loan_id") or raw_dict.get("Loan Id"))
                or effective_loan_id
                or ""
            )
            row_username = rec.get("pkg_username") or rec.get("uploaded_by") or rec.get("username") or pkg.get("uploaded_by") or ""
            row_email = rec.get("pkg_email") or rec.get("uploaded_email") or rec.get("email") or pkg.get("uploaded_email") or ""

            row_obj["Loan Id"] = str(row_loan_id) if row_loan_id is not None else ""
            row_obj["Username"] = str(row_username) if row_username is not None else ""
            row_obj["Email"] = str(row_email) if row_email is not None else ""

            row_obj["Dict Match"] = rec.get("dictionary_match_status") or ""
            row_obj["Type Check"] = rec.get("data_type_validation") or ""
            row_obj["Type Remarks"] = rec.get("data_type_remarks") or ""
            row_obj["Numeric Format"] = rec.get("numeric_format_flag") or ""
            row_obj["Enum Check"] = rec.get("enum_validation") or ""
            row_obj["Enum Remarks"] = rec.get("enum_remarks") or ""
            row_obj["Macro Check"] = rec.get("macro_validation_result") or ""
            row_obj["Combined Result"] = rec.get("result") or ""
            row_obj["Condition Check"] = rec.get("condition_validation") or ""
            row_obj["Condition Remarks"] = rec.get("condition_remarks") or ""
            row_obj["Condition Text"] = rec.get("condition_text") or (raw_dict.get("Condition Text") if isinstance(raw_dict, dict) else "") or ""
            row_obj["Condition Outcome"] = rec.get("condition_outcome") or (raw_dict.get("Condition Outcome") if isinstance(raw_dict, dict) else "") or ""
            row_obj["Final Result"] = rec.get("final_result") or rec.get("result") or ""
            row_obj["Final Remarks"] = rec.get("final_remarks") or ""
            row_obj["Action"] = rec.get("action") or ""
            row_obj["User Remark"] = rec.get("remark") or ""

            ws.append([row_obj.get(h, "") for h in final_headers])

        # Auto-fit column widths (sampling up to 200 rows for fast execution)
        sample_limit = min(len(parsed_records) + 2, 202)
        for col_idx, h in enumerate(final_headers, 1):
            max_len = len(str(h))
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            for row_idx in range(2, sample_limit):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val:
                    val_len = len(str(val))
                    if val_len > max_len:
                        max_len = val_len
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 60), 12)

        output = io.BytesIO()
        wb.save(output)
        content_bytes = output.getvalue()
        output.seek(0)

        date_str = datetime.utcnow().strftime("%Y-%m-%d")
        download_filename = f"Bulk_Audit_Export_{len(packages_rows)}_Packages_{date_str}.xlsx"

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{download_filename}"',
                "Content-Length": str(len(content_bytes)),
                "Access-Control-Expose-Headers": "Content-Disposition, Content-Length"
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Bulk export failed: {str(e)}")

@router.get("/{package_id}")
def get_package_details(package_id: int, user_id: int = Depends(get_current_user_id)):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    try:
        with engine.connect() as conn:
            user_row = conn.execute(
                text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
            ).fetchone()
            if not user_row:
                raise HTTPException(status_code=404, detail="User not found")
            role = user_row[0]

            # Get the package (with uploader info via JOIN)
            lp = conn.execute(
                text("""
                    SELECT lp.*, u.username AS uploaded_by, u.email AS uploaded_email
                    FROM loan_packages lp
                    LEFT JOIN users u ON lp.user_id = u.id
                    WHERE lp.id = :id
                """),
                {"id": package_id}
            ).mappings().first()

            if not lp:
                raise HTTPException(status_code=404, detail="Package not found")

            # Check authorization
            if role != "admin" and lp["user_id"] != user_id:
                raise HTTPException(status_code=403, detail="Not authorized to access this package")

            # Get the records
            records_rows = conn.execute(
                text("SELECT * FROM loan_package_records WHERE package_id = :package_id"),
                {"package_id": package_id}
            ).mappings().all()

            summary = {
                "Total Fields": lp["total_fields"],
                "Matched": lp["matched"],
                "Unmatched": lp["unmatched"],
                "Type Pass": lp["type_pass"],
                "Type Fail": lp["type_fail"],
                "Blank Values": lp["blank_values"],
                "Total Enum": lp["total_enum"],
                "Enum Pass": lp["enum_pass"],
                "Enum Format Variance": lp["enum_format_variance"],
                "Enum Fail": lp["enum_fail"],
                "Enum No List": lp["enum_no_list"],
                "Numeric Format Flagged": lp["numeric_format_flagged"],
                "Total Macro": lp["total_macro"],
                "Macro Valid": lp["macro_valid"],
                "Macro Invalid": lp["macro_invalid"],
                "Final Valid": lp["final_valid"],
                "Final Invalid": lp["final_invalid"],
                "Final Blank": lp["final_blank"]
            }

            records = []
            for r in records_rows:
                rd = dict(r)
                raw_str = rd.get("raw_data")
                if raw_str:
                    try:
                        raw_dict = json.loads(raw_str)
                        if isinstance(raw_dict, dict):
                            raw_dict.pop("condition_text", None)
                            raw_dict.pop("condition_outcome", None)
                            rd["raw_data"] = raw_dict
                            for k, v in raw_dict.items():
                                if k not in rd:
                                    rd[k] = v
                    except Exception:
                        pass
                raw_obj = rd.get("raw_data") if isinstance(rd.get("raw_data"), dict) else {}
                rd["condition_text"] = rd.get("condition_text") or rd.get("Condition Text") or raw_obj.get("Condition Text") or ""
                rd["condition_outcome"] = rd.get("condition_outcome") or rd.get("Condition Outcome") or raw_obj.get("Condition Outcome") or ""
                records.append(rd)

            # Calculate condition validation counts and overall statistics from records
            def _get_val(rec, key1, key2):
                v = rec.get(key1)
                if v is not None and str(v).strip() != "":
                    return str(v).strip()
                raw = rec.get("raw_data")
                if isinstance(raw, dict):
                    v2 = raw.get(key2) or raw.get(key1)
                    if v2 is not None and str(v2).strip() != "":
                        return str(v2).strip()
                return ""

            cond_valid = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "Valid")
            cond_partial = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "Valid - partial")
            cond_invalid = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "Invalid")
            cond_not_auto = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "Not Automated")
            cond_no_rule = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "N/A - no condition rule")
            cond_not_met = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "N/A - condition not met")
            cond_blank = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "N/A - blank value")
            cond_info = sum(1 for r in records if _get_val(r, "condition_validation", "Condition Validation") == "N/A - informational")
            total_rec = len(records)
            cond_known = cond_valid + cond_partial + cond_invalid + cond_not_auto + cond_no_rule + cond_not_met + cond_blank + cond_info
            cond_unclass = total_rec - cond_known

            fr_valid = sum(1 for r in records if _get_val(r, "final_result", "Final Result") == "Valid")
            fr_invalid = sum(1 for r in records if _get_val(r, "final_result", "Final Result") == "Invalid")
            fr_review = sum(1 for r in records if _get_val(r, "final_result", "Final Result") == "Review Required")
            fr_not_validated = sum(1 for r in records if _get_val(r, "final_result", "Final Result") == "Not Validated")

            summary.update({
                "Total Fields": total_rec,
                "Final Valid": fr_valid,
                "Final Invalid": fr_invalid,
                "Review Required": fr_review,
                "Cond Valid": cond_valid,
                "Cond Valid - Partial": cond_partial,
                "Cond Invalid": cond_invalid,
                "Cond Not Automated": cond_not_auto,
                "Cond N/A - No Rule": cond_no_rule,
                "Cond N/A - Condition Not Met": cond_not_met,
                "Cond N/A - Blank Value": cond_blank,
                "Cond N/A - Informational": cond_info,
                "Cond Unclassified": cond_unclass,
                "Cond Rules Applied": cond_valid + cond_partial + cond_invalid + cond_not_auto,
                "Cond Total (reconciles)": total_rec,
                "Overall Valid": fr_valid,
                "Overall Invalid": fr_invalid,
                "Overall Review Required": fr_review,
                "Overall Not Validated": fr_not_validated,
            })

        # Determine loan_id for this package
        lp_dict = dict(lp)
        lp_dict["total_fields"] = total_rec
        lp_dict["final_valid"] = fr_valid
        lp_dict["final_invalid"] = fr_invalid
        lp_dict["review_required"] = fr_review
        loan_id_val = lp_dict.get("loan_id")
        if not loan_id_val:
            for r in records:
                f = str(r.get("data_field") or "").strip().lower()
                if f in ("loan id", "loan_id", "loanid") and r.get("true_data"):
                    loan_id_val = str(r["true_data"]).strip()
                    break

        return {
            "status": "success",
            "package": {
                "id": lp["id"],
                "filename": lp["filename"],
                "loan_id": loan_id_val,
                "uploaded_at": _format_utc_iso(lp["uploaded_at"]) if lp["uploaded_at"] else None,
                "uploaded_by_id": lp["user_id"],
                "uploaded_by": lp.get("uploaded_by"),
                "uploaded_email": lp.get("uploaded_email"),
            },
            "loan_id": loan_id_val,
            "uploaded_by": lp.get("uploaded_by"),
            "uploaded_email": lp.get("uploaded_email"),
            "summary": summary,
            "records": records
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.delete("/clear-all")
def delete_all_packages(user_id: int = Depends(get_current_user_id)):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    with engine.connect() as conn:
        user_row = conn.execute(
            text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
        ).fetchone()
        if not user_row:
            raise HTTPException(status_code=404, detail="User not found")
        role = user_row[0]

    def _do_delete_all():
        with engine.begin() as conn:
            if role == "admin":
                # Acquire table-level exclusive lock on parent first — same order as INSERT
                conn.execute(text("SELECT TOP 0 id FROM loan_packages WITH (TABLOCKX, HOLDLOCK)"))
                conn.execute(text("DELETE FROM loan_package_records"))
                conn.execute(text("DELETE FROM loan_packages"))
            else:
                conn.execute(
                    text("DELETE FROM loan_package_records WHERE package_id IN (SELECT id FROM loan_packages WHERE user_id = :user_id)"),
                    {"user_id": user_id}
                )
                conn.execute(
                    text("DELETE FROM loan_packages WHERE user_id = :user_id"),
                    {"user_id": user_id}
                )

    try:
        _retry_on_deadlock(_do_delete_all)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    return {"status": "success", "message": "All packages deleted successfully"}

@router.delete("/{package_id}")
def delete_package(package_id: int, user_id: int = Depends(get_current_user_id)):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    with engine.connect() as conn:
        user_row = conn.execute(
            text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
        ).fetchone()
        role = user_row[0] if user_row else "user"

    def _do_delete():
        with engine.begin() as conn:
            # UPDLOCK: auth check and delete are atomic — eliminates the TOCTOU gap
            lp = conn.execute(
                text("SELECT user_id FROM loan_packages WITH (UPDLOCK) WHERE id = :id"),
                {"id": package_id}
            ).mappings().first()
            if not lp:
                raise HTTPException(status_code=404, detail="Package not found")
            if role != "admin" and lp["user_id"] != user_id:
                raise HTTPException(status_code=403, detail="Not authorized to delete this package")
            conn.execute(
                text("DELETE FROM loan_package_records WHERE package_id = :package_id"),
                {"package_id": package_id}
            )
            conn.execute(
                text("DELETE FROM loan_packages WHERE id = :id"),
                {"id": package_id}
            )

    try:
        _retry_on_deadlock(_do_delete)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    return {"status": "success", "message": "Package deleted successfully"}

@router.put("/{package_id}/records")
def update_package_records(
    package_id: int,
    payload: Dict[str, Any],
    user_id: int = Depends(get_current_user_id)
):
    if not engine:
        raise HTTPException(status_code=500, detail="Database connection not initialized")

    records = payload.get("records", [])

    with engine.connect() as conn:
        lp = conn.execute(
            text("SELECT user_id FROM loan_packages WHERE id = :id"),
            {"id": package_id}
        ).mappings().first()
        if not lp:
            raise HTTPException(status_code=404, detail="Package not found")
        user_row = conn.execute(
            text("SELECT role FROM users WHERE id = :id"), {"id": user_id}
        ).fetchone()
        role = user_row[0] if user_row else "user"
        if role != "admin" and lp["user_id"] != user_id:
            raise HTTPException(status_code=403, detail="Not authorized to update this package")

    def _do_update():
        with engine.begin() as conn:
            # UPDLOCK prevents S→X upgrade deadlock when concurrent saves are in flight
            lp_check = conn.execute(
                text("SELECT user_id FROM loan_packages WITH (UPDLOCK) WHERE id = :id"),
                {"id": package_id}
            ).mappings().first()
            if not lp_check:
                return

            # Split OR predicate into 3 targeted index seeks — avoids full-table scan locking
            by_id: list = []
            by_field_with_loan: list = []
            by_field_null_loan: list = []
            for r in records:
                rec_id = r.get("id")
                try:
                    rec_id = int(rec_id) if rec_id is not None else None
                except (ValueError, TypeError):
                    rec_id = None
                base = {"package_id": package_id, "action": r.get("action"), "remark": r.get("remark")}
                if rec_id is not None:
                    by_id.append({**base, "record_id": rec_id})
                elif r.get("loan_id") is not None:
                    by_field_with_loan.append({**base, "data_field": r.get("data_field"), "loan_id": r.get("loan_id")})
                else:
                    by_field_null_loan.append({**base, "data_field": r.get("data_field")})

            if by_id:
                conn.execute(
                    text("UPDATE loan_package_records SET action=:action, remark=:remark WHERE package_id=:package_id AND id=:record_id"),
                    by_id
                )
            if by_field_with_loan:
                conn.execute(
                    text("UPDATE loan_package_records SET action=:action, remark=:remark WHERE package_id=:package_id AND data_field=:data_field AND loan_id=:loan_id"),
                    by_field_with_loan
                )
            if by_field_null_loan:
                conn.execute(
                    text("UPDATE loan_package_records SET action=:action, remark=:remark WHERE package_id=:package_id AND data_field=:data_field AND loan_id IS NULL"),
                    by_field_null_loan
                )
            # engine.begin() auto-commits here — UPDLOCK released

    try:
        _retry_on_deadlock(_do_update)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    return {"status": "success", "message": "Package records updated successfully"}
