#!/usr/bin/env python3
"""Audit the condition specification before wiring it into the validator.

Checks, per spec row:
  * does the Data Field exist in the data dictionary?
  * does every clause parse into an automatable rule?
  * for cross-field rules, which fields does it reference, and do those exist?
  * does the Direct / Inter Related label match what the clause actually does?

Writes a workbook with one row per spec row, a column per referenced field,
and a plain statement of what to correct where a clause cannot be automated.

    python audit_conditions.py Condition_Validation_v2.xlsx data_dictionary.xlsx audit.xlsx
"""

import difflib
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

import condition_validation as cv

# Must match validate_loan_files_v5.py, so the audit and the validator can
# never disagree about what is automatable.
REQUIRED_CONDITION_VERSION = "5.5.0"


def _version_tuple(text):
    parts = []
    for chunk in str(text).split("."):
        digits = "".join(c for c in chunk if c.isdigit())
        parts.append(int(digits) if digits else 0)
    return tuple(parts)


def check_parser_version():
    found = getattr(cv, "__version__", None)
    if found is None:
        print(f"ERROR: condition_validation.py declares no __version__; "
              f"expected {REQUIRED_CONDITION_VERSION}. This audit would not "
              f"match the validator. Replace the file.")
        return False
    if _version_tuple(found) < _version_tuple(REQUIRED_CONDITION_VERSION):
        print(f"ERROR: condition_validation.py is {found}, expected "
              f"{REQUIRED_CONDITION_VERSION} or newer. The audit would report "
              f"different results from the validator. Update the file.")
        return False
    return True

SPEC_SHEET_CANDIDATES = ("Data", "Sheet1")
MAX_RELATED_COLUMNS = 6

# Families that read another field's value.
CROSS_FIELD_FAMILIES = {
    "COND_NOT_BLANK", "COND_MUST_BE_BLANK", "COND_EQUALS", "COND_NOT_EQUALS",
    "COND_ALLOWED_SET", "COND_FORMAT",
    "MUTEX", "MATCH_FIELD", "OTHER_FIELD", "DATE_COMPARE", "DATE_OFFSET",
    "DATE_WITHIN", "FIELDS_MATCH_THEN_MATCH", "YEAR_MATCHES_FIELD",
    "DATE_AROUND", "DATE_BETWEEN",
}

UNPARSED = ("NOT_AUTOMATED", "SPEC_MALFORMED")

# Plain-English meaning of each rule family, shown per clause so a reviewer can
# confirm the parser read the sentence the way the business team intended.
FAMILY_NOTES = {
    "NOT_BLANK": "Field must not be blank",
    "NOT_EQUALS": "Field must not equal the listed value(s)",
    "ALWAYS_EQUALS": "Field must always equal the listed value(s)",
    "ALLOWED_SET": "Field must be one of the listed values",
    "VALID_DATE": "Must parse as a date; year cannot start 00/02",
    "US_DATE_FORMAT": "Date must be written month-first (MM/DD/YYYY)",
    "NOT_FUTURE_DATE": "Date must not be in the future",
    "MIN_YEAR": "Year must not precede a minimum",
    "MIN_AGE": "Date of birth must meet a minimum age",
    "FIRST_OF_MONTH": "Date must be the 1st of a month",
    "DATE_COMPARE": "Date must fall before/after another field's date",
    "DATE_OFFSET": "Date must be an exact offset from another field",
    "DATE_WITHIN": "Date must fall within N periods of another field",
    "MUTEX": "If the trigger checkbox is selected, listed options must not be",
    "COND_NOT_BLANK": "If the referenced field matches, this must not be blank",
    "COND_MUST_BE_BLANK": "If the referenced field matches, this must be blank",
    "COND_EQUALS": "If the referenced field matches, this must equal value(s)",
    "COND_ALLOWED_SET": "If the referenced field matches, this must be one of a list",
    "COND_FORMAT": "If the referenced field matches, this must meet a character format",
    "COND_NOT_EQUALS": "If the referenced field matches, this must not equal value(s)",
    "MATCH_FIELD": "Must equal the value of another field",
    "OTHER_FIELD": "Constrains a different field named in the sentence",
    "FIELDS_MATCH_THEN_MATCH": "If two named fields agree, this must equal a third",
    "YEAR_MATCHES_FIELD": "Year must match the year of another field's date",
    "DATE_AROUND": "Date must fall within N periods either side of another field",
    "DATE_BETWEEN": "Date must fall between two other fields' dates",
    "FORMAT": "Character format: length / digits / case / special characters",
    "POSITIVE_NUMBER": "Value must be greater than zero",
    "WHOLE_NUMBER": "Value must have no fractional part",
    "VALUE_RANGE": "Value must fall between two bounds",
    "PERCENTAGE_FORMAT": "Value must be a percentage between 0 and 100",
    "DECIMAL_SET": "Fractional part must be one of a permitted set",
    "GREATER_THAN": "Numeric value must exceed a threshold",
    "MIN_VALUE": "Numeric value must be at least a threshold",
    "MAX_VALUE": "Numeric value must not exceed a threshold",
    "EXACT_DIGITS": "Exact digit count",
    "INFORMATIONAL": "Descriptive note; no testable constraint",
    "SPEC_MALFORMED": "Spec sentence cannot be read (unbalanced quotes)",
    "NOT_AUTOMATED": "Not interpreted - manual review required",
}


def norm(text):
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


def load_spec(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = next((workbook[n] for n in SPEC_SHEET_CANDIDATES
                  if n in workbook.sheetnames), workbook.worksheets[0])
    header = {}
    for cell in sheet[1]:
        if cell.value is not None:
            header[" ".join(str(cell.value).split())] = cell.column

    rows = []
    for index in range(2, sheet.max_row + 1):
        def get(name):
            column = header.get(name)
            return sheet.cell(row=index, column=column).value if column else None
        rows.append({
            "S.No": get("S.No"),
            "Data Field": " ".join(str(get("Data Field") or "").split()),
            "True Source": " ".join(str(get("True Source") or "").split()),
            "Spec Kind": " ".join(str(
                get("Direct Field/Inter Related") or "").split()),
            "Conditions": get("Conditions"),
        })
    return rows


def referenced_fields(rule):
    """Field names a parsed rule reads, besides the row's own field."""
    names = []
    params = rule.params or {}

    if rule.family == "MUTEX":
        # The trigger is a sibling checkbox within this field's own group.
        option = params.get("trigger_option")
        if option:
            names.append(f"(same group) {option}")
    else:
        if params.get("ref_field"):
            names.append(params["ref_field"])
        for extra in params.get("extra_conditions") or []:
            if extra.get("ref_field"):
                names.append(extra["ref_field"])
        if params.get("subject"):
            names.append(params["subject"])
        # Both sides of a two-field comparison are referenced fields.
        for side in ("left", "right"):
            if params.get(side):
                names.append(params[side])

    guard = params.get("guard") or {}
    if guard.get("ref_field"):
        names.append(guard["ref_field"])
    for extra in guard.get("extra_conditions") or []:
        if extra.get("ref_field"):
            names.append(extra["ref_field"])

    seen, out = set(), []
    for name in names:
        key = norm(name)
        if key and key not in seen:
            seen.add(key)
            out.append(name)
    return out


def resolve_in_dictionary(reference, dictionary_norm, dictionary_names):
    """Mirror the validator's lookup so the audit reflects runtime behaviour."""
    if reference.startswith("(same group) "):
        return "sibling checkbox", "OK"

    key = norm(reference)
    if key in dictionary_norm:
        return dictionary_norm[key], "OK"

    bare = norm(re.sub(r"\([^)]*\)", " ", reference))
    if bare and bare in dictionary_norm:
        return dictionary_norm[bare], "OK (parenthetical dropped)"

    # Borrower-scoped references resolve at runtime against the row's own
    # borrower, so several same-named candidates are expected, not a defect.
    ends = [n for n in dictionary_names if norm(n).endswith(key)]
    if len(ends) == 1:
        return ends[0], "OK (suffix match)"
    if len(ends) > 1:
        return "; ".join(sorted(ends)[:3]), "AMBIGUOUS"

    contains = [n for n in dictionary_names if key and key in norm(n)]
    if len(contains) == 1:
        return contains[0], "OK (partial match)"
    if len(contains) > 1:
        return "; ".join(sorted(contains)[:3]), "AMBIGUOUS"
    return "", "NOT IN DICTIONARY"


# Clauses whose meaning is clear and unambiguous, but whose wording the parser
# does not yet accept. These need a CODE change, not a spec rewrite -- telling
# the business team to reword correct English would be wrong.
PARSER_GAPS = [
    (r"(?i)must contain exactly \d+ numeric digits", "exact digit count"),
    (r"(?i)must be a positive\s*(number|#)", "positive number"),
    (r'(?i)(greater|less)\s+than\s+the\s+"', "date/value compared to another field"),
    (r"(?i)must be exactly one year after", "date offset from another field"),
    (r"(?i)age must be greater than or equal to \d+", "minimum age"),
    (r"(?i)decimal portion of the value must be one of", "permitted decimals"),
    (r"(?i)must contain only lowercase", "lowercase only"),
    (r"(?i)must contain a whole number", "whole number"),
    (r"(?i)must contain (only )?numeric characters only|must contain a valid numeric",
     "numeric only"),
    (r"(?i)(in|be always) percentage format", "percentage format"),
    (r"(?i)should be the 1st day of a month", "first day of month"),
    (r"(?i)must contain a decimal value between", "decimal range"),
    (r'(?i)same as be "|should (be )?same as "', "match another field"),
    (r"(?i)not checked then|screen checked then", "checkbox-driven condition"),
]


def parser_gap(text):
    for pattern, label in PARSER_GAPS:
        if re.search(pattern, text):
            return label
    return None


NUMERIC_FAMILIES = {"MAX_VALUE", "MIN_VALUE", "GREATER_THAN", "EXACT_DIGITS",
                    "DECIMAL_SET"}
DATE_WORDS = re.compile(r"(?i)\b(date|day|week|month|year)s?\b")


def suspicious_parse(rule, enum_values=None):
    """Flag parses that look wrong, so a bad automation is not mistaken for a
    good one. A wrong verdict is worse than no verdict."""
    params = rule.params or {}
    text = rule.text
    notes = []

    if rule.family in NUMERIC_FAMILIES and DATE_WORDS.search(text):
        notes.append(f"{rule.family} assigned to a clause mentioning dates/"
                     f"periods - verify this is really a numeric bound")

    literal = str(params.get("literal") or "")
    if literal and re.search(r"[0-9]\s*[a-z]", literal, re.I):
        notes.append(f"threshold parsed from {literal!r} - a unit word may "
                     f"have been absorbed")

    for key in ("ref_values", "forbidden", "allowed", "expected"):
        for value in params.get(key) or []:
            if re.search(r"(?i)\b(screen|field)\b", str(value)):
                notes.append(f"{key} contains {value!r}, which looks like a "
                             f"field or screen name, not a value")
            # Only suspect a date-ish token when it is NOT a permitted value
            # for this field. "45 Days" is a legitimate enumeration member of
            # Index Lookback Period, not a mis-read field name.
            if (DATE_WORDS.search(str(value))
                    and rule.family in ("NOT_EQUALS", "ALWAYS_EQUALS",
                                        "ALLOWED_SET")
                    and norm(value) not in (enum_values or set())):
                notes.append(f"{key} contains {value!r} - a date field name "
                             f"may have been read as a literal value")

    reference = params.get("ref_field")
    if reference and re.fullmatch(r"(?i)\s*(note|1003|title|property|"
                                  r"employment-income)\s*", str(reference)):
        notes.append(f"referenced field {reference!r} looks like a screen name")

    return " ; ".join(dict.fromkeys(notes))


def correction_for(rule):
    """Plain statement of what to fix so a clause can be automated."""
    text = rule.text

    if rule.family == "SPEC_MALFORMED":
        return ("Unbalanced quotation marks",
                "Close every quoted value. The value list cannot be read "
                "while a quote is left open.")

    quoted_tokens = re.findall(r'"([^"]*)"', text)
    padded = [t for t in quoted_tokens if t != t.strip() and t.strip()]

    checks = [
        (None if not padded else "MATCH",
         "Leading/trailing space inside quotes",
         'Remove the space inside the quotes (e.g. " Condo" -> "Condo"); the '
         'space becomes part of the value and never matches.'),
        (r"(?i)should be applicable|can be a value of|this may be|"
         r"as applicable",
         "Vague consequent",
         'Replace with an explicit outcome, e.g. "then this should never be '
         '\\"Blank\\"" or "then this should be \\"<value>\\"".'),
        (r"(?i)remaining cases|then,\s*this will be",
         "Two rules in one sentence",
         "Split into two numbered clauses, one per branch, using "
         '"states \\"X\\"" and "states other than \\"X\\"".'),
        (r"(?i)\bmuste\b|fieldnever|\bshould\s+\"|\bthis should \"",
         "Typo or missing verb",
         'Fix the wording so the clause reads "... then this should be '
         '\\"<value>\\"" or "... then this should never be \\"Blank\\"".'),
        (r"(?i)at least \d+ years old|year from the|after whole number",
         "Derived value",
         "Supported now; if still unparsed, state the rule explicitly "
         "(e.g. \"Must be always 1 year from the \\\"<field>\\\" field\")."),
        (r"(?i)\(ex[:\s]|multiple numbers",
         "Rule given by example only",
         'State the rule itself, e.g. "Value must contain at least 3 digits" '
         'or "Value must contain only numeric digits".'),
        (r"(?i)format\s*\d+[-/]\d+.*no special character",
         "Self-contradictory",
         "The sentence prescribes a punctuated format and also forbids "
         "punctuation. Keep one."),
        (r"(?i)should always match with this",
         "Subject reversed",
         'Rewrite so this row is the subject: "This should always match the '
         '\\"<field>\\" field".'),
    ]
    for pattern, issue, fix in checks:
        if pattern is None:
            continue
        if pattern == "MATCH" or re.search(pattern, text):
            return issue, fix

    gap = parser_gap(text)
    if gap:
        return (f"Parser gap - {gap}",
                "NO SPEC CHANGE NEEDED. The clause is clear; the validator "
                "does not yet recognise this wording. Add support in "
                "condition_validation.py.")

    return ("Sentence structure not recognised",
            'Rewrite in a supported shape: "This should never be \\"Blank\\"", '
            '"This should be always \\"<value>\\"", "It should be either '
            '\\"A\\" or \\"B\\"", or "If \\"<field>\\" ... states \\"<value>\\" '
            'then this should be \\"<value>\\"".')


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        return 1
    spec_path, dict_path, out_path = (Path(sys.argv[1]), Path(sys.argv[2]),
                                      Path(sys.argv[3]))

    if not check_parser_version() and "--ignore-version" not in sys.argv:
        print("Refusing to run. Pass --ignore-version to override.")
        return 1
    print(f"Condition parser version "
          f"{getattr(cv, '__version__', 'unknown')}")

    dictionary = pd.read_excel(dict_path, keep_default_na=False, na_values=[""])
    dictionary_names = [str(x).strip() for x in dictionary["Field Name"]
                        if str(x).strip()]
    dictionary_norm = {}
    for name in dictionary_names:
        dictionary_norm.setdefault(norm(name), name)

    # Permitted enumeration values per field, so the suspicious-parse check can
    # tell a real value from a mis-read field name.
    enum_lookup = {}
    if "Field Enumerations" in dictionary.columns:
        for _, row in dictionary.iterrows():
            field_name = str(row.get("Field Name", "")).strip()
            raw = row.get("Field Enumerations", "")
            if not field_name or not str(raw).strip():
                continue
            values = {norm(v) for v in re.split(r"[\r\n]+", str(raw))
                      if str(v).strip()}
            enum_lookup.setdefault(norm(field_name), set()).update(values)

    spec_rows = load_spec(spec_path)
    records = []
    clause_records = []

    for row in spec_rows:
        field = row["Data Field"]
        field_status = ("OK" if norm(field) in dictionary_norm
                        else "NOT IN DICTIONARY")
        field_suggestion = ""
        if field_status != "OK":
            close = difflib.get_close_matches(field, dictionary_names, n=2,
                                              cutoff=0.6)
            field_suggestion = ("did you mean "
                                + " or ".join(repr(c) for c in close) + "?"
                                ) if close else ("no similar field in the "
                                                 "dictionary")

        clauses = cv._split_clauses(row["Conditions"])
        parsed, guard = [], None
        for clause in clauses:
            rule = cv.parse_clause(cv.Rule(
                field=field, source=row["True Source"],
                kind=row["Spec Kind"], text=clause))
            # Mirror load_rules: a bare clause inherits the preceding condition.
            if rule.family in cv.CONDITIONAL_FAMILIES:
                guard = {
                    "ref_field": rule.params.get("ref_field"),
                    "ref_values": rule.params.get("ref_values"),
                    "starts_with": rule.params.get("starts_with", False),
                    "negate": rule.params.get("negate", False),
                    "extra_conditions": rule.params.get("extra_conditions") or [],
                }
            elif (guard and rule.family in cv.GUARDABLE_FAMILIES
                  and not re.search(r"(?i)\balways\b", rule.text)):
                rule.params = dict(rule.params)
                rule.params["guard"] = guard
            parsed.append(rule)

        automated = [r for r in parsed if r.family not in UNPARSED]
        failed = [r for r in parsed if r.family in UNPARSED]

        if not clauses:
            coverage = "No clauses"
        elif not failed:
            coverage = "Fully automatable"
        elif automated:
            coverage = "Partially automatable"
        else:
            coverage = "Not automatable"

        # Referenced fields across all clauses for this row.
        references = []
        for rule in parsed:
            for reference in referenced_fields(rule):
                if reference not in references:
                    references.append(reference)

        resolved, problems, suggestions = [], [], []
        for reference in references:
            match, status = resolve_in_dictionary(
                reference, dictionary_norm, dictionary_names)
            resolved.append((reference, match, status))
            if status == "NOT IN DICTIONARY":
                close = difflib.get_close_matches(
                    reference, dictionary_names, n=2, cutoff=0.6)
                problems.append(f"{reference!r}: not in dictionary")
                if close:
                    suggestions.append(
                        f"{reference!r} -> did you mean "
                        + " or ".join(repr(c) for c in close) + "?")
                else:
                    suggestions.append(
                        f"{reference!r}: no similar field exists; confirm the "
                        f"field name with the business team.")
            elif status == "AMBIGUOUS":
                problems.append(f"{reference!r}: matches several fields")
                suggestions.append(
                    f"{reference!r} matches {match}. Qualify the name in the "
                    f"clause (e.g. include the borrower or category) if the "
                    f"intended field is not the row's own borrower.")

        is_cross = any(r.family in CROSS_FIELD_FAMILIES for r in parsed)
        label = row["Spec Kind"].lower()
        if label.startswith("direct") and is_cross:
            label_check = "MISMATCH - labelled Direct but reads another field"
        elif label.startswith("inter") or label.startswith("related"):
            label_check = ("OK" if is_cross or failed
                           else "MISMATCH - labelled Related but self-contained")
        else:
            label_check = "OK"

        suspicious = []
        for rule in parsed:
            note = suspicious_parse(rule, enum_lookup.get(norm(field), set()))
            if note:
                suspicious.append(f"[{rule.family}] {note}")

        issues, fixes = [], []
        for rule in failed:
            issue, fix = correction_for(rule)
            issues.append(f"[clause] {rule.text[:90]}")
            fixes.append(f"{issue}: {fix}")

        # One record per clause, with its own verdict. A clause is clean only
        # if it parses AND every field it references resolves.
        for position, rule in enumerate(parsed, start=1):
            clause_refs = referenced_fields(rule)
            clause_problems = []
            clause_resolved = []
            for reference in clause_refs:
                match, status = resolve_in_dictionary(
                    reference, dictionary_norm, dictionary_names)
                clause_resolved.append(f"{reference} -> {status}")
                if not status.startswith("OK"):
                    clause_problems.append(f"{reference!r}: {status}")

            unparsed = rule.family in UNPARSED
            gap = parser_gap(rule.text) if unparsed else None
            suspicious = suspicious_parse(
                rule, enum_lookup.get(norm(field), set()))

            if unparsed:
                clause_issue, clause_fix = correction_for(rule)
                owner = "Code (parser gap)" if gap else "Business (spec)"
                clause_verdict = "NEEDS FIX - cannot parse"
            elif clause_problems:
                clause_issue = "; ".join(clause_problems)
                clause_fix = ("Correct the referenced field name, or confirm "
                              "which dictionary field is meant.")
                owner = "Business (field name)"
                clause_verdict = "NEEDS FIX - reference unresolved"
            elif suspicious:
                clause_issue, clause_fix = suspicious, "Verify the parse."
                owner = "Review"
                clause_verdict = "NEEDS REVIEW - parse may be wrong"
            elif field_status != "OK":
                clause_issue = "Parent Data Field not in dictionary"
                clause_fix = field_suggestion
                owner = "Business (field name)"
                clause_verdict = "NEEDS FIX - parent field missing"
            else:
                clause_issue = clause_fix = owner = ""
                clause_verdict = "OK - no issues"

            clause_records.append({
                "S.No": row["S.No"],
                "Data Field": field,
                "Spec Kind": row["Spec Kind"],
                "Clause #": position,
                "Clause": rule.text,
                "Rule Family": rule.family,
                "Interpretation": FAMILY_NOTES.get(rule.family, ""),
                "Cross-Field": "Yes" if rule.family in CROSS_FIELD_FAMILIES
                               else "No",
                "Referenced Fields": "; ".join(clause_refs),
                "Reference Status": "; ".join(clause_resolved),
                "Clause Verdict": clause_verdict,
                "Owner": owner,
                "Issue": clause_issue,
                "Fix": clause_fix,
            })

        record = {
            "S.No": row["S.No"],
            "Data Field": field,
            "Field In Dictionary": field_status,
            "Field Name Suggestion": field_suggestion,
            "True Source": row["True Source"],
            "Spec Kind": row["Spec Kind"],
            "Label Check": label_check,
            "Clauses": len(clauses),
            "Automatable Clauses": len(automated),
            "Coverage": coverage,
            "Rule Families": ", ".join(r.family for r in parsed),
            "Cross-Field": "Yes" if is_cross else "No",
            "Related Field Count": len(references),
            "Related Field Issues": " | ".join(problems),
            "Related Field Suggestion": " || ".join(suggestions),
            "Problem Clauses": " || ".join(issues),
            "Correction Needed": " || ".join(fixes),
            "Suspicious Parse": " || ".join(suspicious),
            "Conditions": " || ".join(clauses),
        }
        for slot in range(MAX_RELATED_COLUMNS):
            if slot < len(resolved):
                reference, match, status = resolved[slot]
                record[f"Related Field {slot + 1}"] = reference
                record[f"Related Field {slot + 1} Status"] = status
            else:
                record[f"Related Field {slot + 1}"] = ""
                record[f"Related Field {slot + 1} Status"] = ""
        if len(resolved) > MAX_RELATED_COLUMNS:
            record["Related Field Issues"] = (
                f"{len(resolved)} referenced fields (only "
                f"{MAX_RELATED_COLUMNS} columns shown) | "
                + record["Related Field Issues"])
        records.append(record)

    clause_frame = pd.DataFrame(clause_records)
    frame = pd.DataFrame(records)

    # A single verdict per field, combining EVERY blocker. A row that parses
    # but references a field that does not exist would fail at runtime, so
    # clause coverage alone is not a clean bill of health.
    def verdict(row):
        blockers = []
        if row["Field In Dictionary"] != "OK":
            blockers.append("Data Field not in dictionary")
        if row["Coverage"] == "Not automatable":
            blockers.append("no clause can be automated")
        elif row["Coverage"] == "Partially automatable":
            blockers.append("some clauses cannot be automated")
        if str(row["Related Field Issues"] or ""):
            blockers.append("referenced field cannot be resolved")
        if str(row["Suspicious Parse"] or ""):
            blockers.append("parse looks wrong - verify")

        if not blockers:
            return "READY - no issues", ""
        if row["Coverage"] == "Not automatable":
            return "NOT AUTOMATABLE", "; ".join(blockers)
        return "NEEDS FIX", "; ".join(blockers)

    verdicts = frame.apply(verdict, axis=1, result_type="expand")
    frame["Verdict"] = verdicts[0]
    frame["Blockers"] = verdicts[1]

    needs_work = frame[frame["Verdict"] != "READY - no issues"].sort_values(
        ["Verdict", "Spec Kind", "S.No"])

    # Verdict x Direct/Related, which is the view that drives the fix list.
    breakdown = (frame.groupby(["Verdict", "Spec Kind"]).size()
                 .unstack(fill_value=0))
    breakdown["Total"] = breakdown.sum(axis=1)
    breakdown = breakdown.reset_index()

    blocker_counts = {}
    for row in frame[frame["Verdict"] != "READY - no issues"].itertuples():
        for blocker in str(row.Blockers).split("; "):
            if blocker:
                blocker_counts[blocker] = blocker_counts.get(blocker, 0) + 1
    blockers_frame = (pd.DataFrame(
        [{"Blocker": k, "Fields Affected": v}
         for k, v in blocker_counts.items()])
        .sort_values("Fields Affected", ascending=False))

    clause_summary = (clause_frame.groupby(["Clause Verdict", "Spec Kind"])
                      .size().unstack(fill_value=0))
    clause_summary["Total"] = clause_summary.sum(axis=1)
    clause_summary = clause_summary.reset_index()

    clause_owner = (clause_frame[clause_frame["Owner"] != ""]
                    .groupby("Owner").size().reset_index(name="Clauses")
                    .sort_values("Clauses", ascending=False))

    summary = pd.DataFrame([
        {"Metric": "Spec rows (fields)", "Count": len(frame)},
        {"Metric": "  of which Direct",
         "Count": int((frame["Spec Kind"] == "Direct").sum())},
        {"Metric": "  of which Related",
         "Count": int((frame["Spec Kind"] == "Related").sum())},
        {"Metric": "READY - fully automatable, no issues",
         "Count": int((frame["Verdict"] == "READY - no issues").sum())},
        {"Metric": "NEEDS FIX - partially blocked",
         "Count": int((frame["Verdict"] == "NEEDS FIX").sum())},
        {"Metric": "NOT AUTOMATABLE",
         "Count": int((frame["Verdict"] == "NOT AUTOMATABLE").sum())},
        {"Metric": "Data Field found in dictionary",
         "Count": (frame["Field In Dictionary"] == "OK").sum()},
        {"Metric": "Data Field NOT in dictionary",
         "Count": (frame["Field In Dictionary"] != "OK").sum()},
        {"Metric": "Fully automatable",
         "Count": (frame["Coverage"] == "Fully automatable").sum()},
        {"Metric": "Partially automatable",
         "Count": (frame["Coverage"] == "Partially automatable").sum()},
        {"Metric": "Not automatable",
         "Count": (frame["Coverage"] == "Not automatable").sum()},
        {"Metric": "Cross-field rows", "Count": (frame["Cross-Field"] == "Yes").sum()},
        {"Metric": "Rows with a referenced-field problem",
         "Count": (frame["Related Field Issues"] != "").sum()},
        {"Metric": "Direct/Related label mismatches",
         "Count": (frame["Label Check"] != "OK").sum()},
        {"Metric": "Rows needing a SPEC correction",
         "Count": int((frame["Correction Needed"].fillna("").str.contains(
             "Parser gap") == False)[frame["Coverage"] != "Fully automatable"].sum())},
        {"Metric": "Rows blocked only by a PARSER gap (no spec change)",
         "Count": int(frame["Correction Needed"].fillna("").str.contains(
             "Parser gap").sum())},
        {"Metric": "Rows with a SUSPICIOUS parse (wrong automation risk)",
         "Count": int((frame["Suspicious Parse"].fillna("") != "").sum())},
        {"Metric": "--- CLAUSE LEVEL ---", "Count": None},
        {"Metric": "Total clauses", "Count": len(clause_frame)},
        {"Metric": "Clauses OK - no issues",
         "Count": int((clause_frame["Clause Verdict"] == "OK - no issues").sum())},
        {"Metric": "Clauses needing a fix",
         "Count": int((clause_frame["Clause Verdict"].str.startswith(
             "NEEDS FIX")).sum())},
        {"Metric": "Clauses needing review",
         "Count": int((clause_frame["Clause Verdict"].str.startswith(
             "NEEDS REVIEW")).sum())},
        {"Metric": "Automatable clauses",
         "Count": int(frame["Automatable Clauses"].sum())},
    ])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path) as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        breakdown.to_excel(writer, sheet_name="Verdict by Type", index=False)
        blockers_frame.to_excel(writer, sheet_name="Blockers", index=False)
        clause_summary.to_excel(writer, sheet_name="Clause Summary",
                                index=False)
        clause_owner.to_excel(writer, sheet_name="Clause Fix Owner",
                              index=False)
        clause_frame[clause_frame["Clause Verdict"] != "OK - no issues"].to_excel(
            writer, sheet_name="Clauses Needing Fix", index=False)
        clause_frame.to_excel(writer, sheet_name="All Clauses", index=False)
        needs_work.to_excel(writer, sheet_name="Needs Correction", index=False)
        frame.to_excel(writer, sheet_name="All Rows", index=False)
        frame[frame["Suspicious Parse"] != ""].to_excel(
            writer, sheet_name="Suspicious Parses", index=False)
        frame[frame["Cross-Field"] == "Yes"].to_excel(
            writer, sheet_name="Cross-Field Map", index=False)

    print(summary.to_string(index=False))
    print()
    print(breakdown.to_string(index=False))
    print()
    print(blockers_frame.to_string(index=False))
    print()
    print(clause_summary.to_string(index=False))
    print()
    print(clause_owner.to_string(index=False))
    print(f"\nWritten: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
